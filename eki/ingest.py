"""Bounded, provenance-preserving parsing. Never execute uploaded content."""
import csv
import hashlib
import io
import re
import socket
import ipaddress
import zipfile
from pathlib import PurePath
from urllib.parse import urlsplit

import httpx
import tiktoken
from bs4 import BeautifulSoup
from docx import Document as WordDocument
from pypdf import PdfReader
from .models import Document, Chunk

ENCODER = tiktoken.get_encoding("cl100k_base")


def parse_file(name: str, data: bytes, max_chars: int = 1_000_000) -> list[Document]:
    name = PurePath(name.replace("\\", "/")).name
    suffix = PurePath(name).suffix.lower()
    docs = []
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted or len(reader.pages) > 300:
            raise ValueError("Encrypted PDFs and PDFs over 300 pages are unsupported")
        for n, page in enumerate(reader.pages, 1):
            docs.append(Document(content=page.extract_text() or "", source=name, page=n, document_type="pdf"))
    elif suffix == ".docx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 20_000_000:
                raise ValueError("Expanded DOCX is too large")
        doc = WordDocument(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        parts += [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        docs.append(Document(content="\n\n".join(parts), source=name, document_type="docx"))
    elif suffix == ".csv":
        reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
        if not reader.fieldnames or len(reader.fieldnames) > 100:
            raise ValueError("CSV requires a header with at most 100 columns")
        for rownum, row in enumerate(reader, 2):
            if rownum > 10001:
                raise ValueError("CSV exceeds 10,000 rows")
            docs.append(Document(content="; ".join(f"{k}: {v}" for k, v in row.items()),
                                 source=name, row=rownum, document_type="csv"))
    elif suffix in (".html", ".htm", ".txt", ".md"):
        text = data.decode("utf-8-sig")
        if suffix in (".html", ".htm"):
            soup = BeautifulSoup(text, "html.parser")
            for tag in soup(["script", "style", "noscript", "nav"]):
                tag.decompose()
            text = soup.get_text("\n", strip=True)
        docs.append(Document(content=text, source=name, document_type=suffix[1:]))
    else:
        raise ValueError("Supported formats: PDF, DOCX, CSV, HTML, TXT, MD")
    if sum(len(d.content) for d in docs) > max_chars:
        raise ValueError("Extracted text exceeds limit")
    docs = [d for d in docs if d.content.strip()]
    if not docs:
        raise ValueError("No extractable text; scanned PDFs need external OCR")
    return docs


def fetch_web(url: str, allowed_hosts: list[str], max_bytes: int) -> bytes:
    parsed = urlsplit(url)
    if (parsed.scheme != "https" or parsed.hostname not in allowed_hosts or parsed.username
            or parsed.password or parsed.port not in (None, 443)):
        raise ValueError("Only explicitly allowlisted HTTPS hosts on port 443 are permitted")
    addresses = socket.getaddrinfo(parsed.hostname, 443, type=socket.SOCK_STREAM)
    if not addresses or any(not ipaddress.ip_address(a[4][0]).is_global for a in addresses):
        raise ValueError("Private, loopback and link-local addresses are forbidden")
    # Host allowlist is admin controlled. Production must ALSO enforce egress at the network layer
    # to prevent DNS rebinding between this resolution and httpx's connection resolution.
    with httpx.Client(timeout=15, follow_redirects=False, trust_env=False) as client:
        with client.stream("GET", url) as response:
            if response.status_code != 200:
                raise ValueError("Web fetch requires a direct HTTP 200 response")
            if "text/html" not in response.headers.get("content-type", ""):
                raise ValueError("Web loader accepts HTML only")
            chunks, size = [], 0
            for block in response.iter_bytes():
                size += len(block)
                if size > max_bytes:
                    raise ValueError("Web document too large")
                chunks.append(block)
            return b"".join(chunks)


def chunk_documents(docs: list[Document], document_id: str, strategy: str = "recursive",
                    max_tokens: int = 220, embedder=None) -> list[Chunk]:
    if strategy not in ("recursive", "token", "semantic"):
        raise ValueError("Unknown chunk strategy")
    if max_tokens < 16:
        raise ValueError("max_tokens must be at least 16")
    result = []
    for doc in docs:
        if strategy == "token":
            tokens = ENCODER.encode(doc.content)
            pieces = [ENCODER.decode(tokens[i:i + max_tokens]) for i in range(0, len(tokens), max_tokens)]
        else:
            sentences = [s.strip() for s in re.split(r"\n\s*\n|(?<=[.!?])\s+", doc.content) if s.strip()]
            units = []
            for sentence in sentences:
                tokens = ENCODER.encode(sentence)
                units.extend(ENCODER.decode(tokens[i:i + max_tokens]) for i in range(0, len(tokens), max_tokens))
            vectors = None
            if strategy == "semantic":
                if embedder is None or not embedder.semantic:
                    raise ValueError("Semantic chunking requires learned embeddings (fastembed)")
                vectors = embedder.encode(units)
            pieces, current, count = [], [], 0
            for i, unit in enumerate(units):
                size = len(ENCODER.encode(unit))
                boundary = vectors is not None and i > 0 and float(vectors[i] @ vectors[i - 1]) < 0.55
                if current and (count + size + 1 > max_tokens or boundary):
                    pieces.append(" ".join(current))
                    current, count = [], 0
                current.append(unit)
                count += size + 1
            if current:
                pieces.append(" ".join(current))
        for piece in pieces:
            ordinal = len(result)
            cid = hashlib.sha256(f"{document_id}:{ordinal}:{piece}".encode()).hexdigest()[:32]
            result.append(Chunk(**doc.model_dump(exclude={"content"}), content=piece, id=cid,
                                document_id=document_id, ordinal=ordinal))
    return result
